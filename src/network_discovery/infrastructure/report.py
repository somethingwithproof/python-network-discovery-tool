"""Report service implementations.

This module provides implementations of the ReportService interface
with unified export capabilities for JSON, CSV, PDF, and other formats.
"""

from __future__ import annotations

import csv
import importlib.util
import json
import logging
from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING, Any

import jinja2
import openpyxl

from network_discovery.application.interfaces import ReportService


if TYPE_CHECKING:
    from network_discovery.domain.device import Device


# Setup logging
logger = logging.getLogger(__name__)

# Check for optional PDF dependencies
FPDF_AVAILABLE = importlib.util.find_spec("fpdf2") is not None
if not FPDF_AVAILABLE:
    logger.debug("fpdf2 not available. PDF export will be disabled.")


class ReportGenerator(ReportService):
    """Implementation of ReportService for generating reports in various formats."""

    def __init__(self, output_dir: str, template_dir: str | None = None) -> None:
        """Initialize a new ReportGenerator.

        Args:
            output_dir: The directory where reports will be saved.
            template_dir: The directory containing HTML templates.
        """
        self.output_dir = output_dir
        self.template_dir = template_dir
        Path(output_dir).mkdir(parents=True, exist_ok=True)

    @property
    def supported_formats(self) -> list[str]:
        """Return list of supported export formats.

        Returns:
            List of format strings that can be used with generate_report.
        """
        formats = ["html", "csv", "xlsx", "json"]
        if FPDF_AVAILABLE:
            formats.append("pdf")
        return formats

    def generate_report(
        self,
        devices: list[Device],
        format_type: str,
        filename: str | None = None,
    ) -> Any:
        """Generate a report for a list of devices.

        Args:
            devices: The list of devices to include in the report.
            format_type: The format of the report (e.g., "html", "csv", "xlsx", "json", "pdf").
            filename: Optional custom filename (without extension).

        Returns:
            The path to the generated report file.

        Raises:
            ValueError: If the format type is not supported.
        """
        format_type = format_type.lower()

        generators = {
            "html": self._generate_html_report,
            "csv": self._generate_csv_report,
            "xlsx": self._generate_excel_report,
            "json": self._generate_json_report,
            "pdf": self._generate_pdf_report,
        }

        generator_fn = generators.get(format_type)
        if generator_fn is None:
            raise ValueError(
                f"Unsupported report format: {format_type}. "
                f"Supported formats: {', '.join(self.supported_formats)}"
            )
        return generator_fn(devices, filename)

    def _get_output_path(self, extension: str, filename: str | None = None) -> str:
        """Generate output file path with timestamp if no filename provided.

        Args:
            extension: File extension (e.g., "csv", "json", "pdf").
            filename: Optional custom filename (without extension).

        Returns:
            Full path to the output file.
        """
        if filename:
            base_name = filename
        else:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            base_name = f"devices_{timestamp}"
        return str(Path(self.output_dir) / f"{base_name}.{extension}")

    def _generate_html_report(
        self, devices: list[Device], filename: str | None = None
    ) -> str:
        """Generate an HTML report.

        Args:
            devices: The list of devices to include in the report.
            filename: Optional custom filename (without extension).

        Returns:
            The path to the generated HTML file.
        """
        if not self.template_dir:
            raise ValueError("Template directory is required for HTML reports")

        env = jinja2.Environment(
            loader=jinja2.FileSystemLoader(self.template_dir),
            autoescape=jinja2.select_autoescape(["html", "xml"]),
        )
        template = env.get_template("layout.html")
        output = template.render(
            devices=devices,
            generated_at=datetime.now().isoformat(),
            total_devices=len(devices),
            alive_devices=sum(1 for d in devices if d.alive),
        )

        output_path = self._get_output_path("html", filename)
        with Path(output_path).open("w", encoding="utf-8") as f:
            f.write(output)

        logger.info("HTML report generated at %s", output_path)
        return output_path

    def _generate_csv_report(
        self, devices: list[Device], filename: str | None = None
    ) -> str:
        """Generate a CSV report.

        Args:
            devices: The list of devices to include in the report.
            filename: Optional custom filename (without extension).

        Returns:
            The path to the generated CSV file.
        """
        output_path = self._get_output_path("csv", filename)
        with Path(output_path).open("w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(
                [
                    "ID",
                    "Host",
                    "IP",
                    "SNMP Group",
                    "Alive",
                    "SNMP",
                    "SSH",
                    "MySQL",
                    "Uname",
                    "Scanned",
                    "Errors",
                ]
            )
            for device in devices:
                writer.writerow(
                    [
                        device.id,
                        device.host,
                        device.ip,
                        device.snmp_group,
                        device.alive,
                        device.snmp,
                        device.ssh,
                        device.mysql,
                        device.uname,
                        device.scanned,
                        "; ".join(device.errors),
                    ]
                )

        logger.info("CSV report generated at %s", output_path)
        return output_path

    def _generate_excel_report(
        self, devices: list[Device], filename: str | None = None
    ) -> str:
        """Generate an Excel report.

        Args:
            devices: The list of devices to include in the report.
            filename: Optional custom filename (without extension).

        Returns:
            The path to the generated Excel file.
        """
        output_path = self._get_output_path("xlsx", filename)
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Devices"

        # Write header
        headers = [
            "ID",
            "Host",
            "IP",
            "SNMP Group",
            "Alive",
            "SNMP",
            "SSH",
            "MySQL",
            "Uname",
            "Scanned",
            "Errors",
        ]
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col, value=header)
            cell.font = openpyxl.styles.Font(bold=True)

        # Write data
        for row, device in enumerate(devices, 2):
            sheet.cell(row=row, column=1, value=device.id)
            sheet.cell(row=row, column=2, value=device.host)
            sheet.cell(row=row, column=3, value=device.ip)
            sheet.cell(row=row, column=4, value=device.snmp_group)
            sheet.cell(row=row, column=5, value=device.alive)
            sheet.cell(row=row, column=6, value=device.snmp)
            sheet.cell(row=row, column=7, value=device.ssh)
            sheet.cell(row=row, column=8, value=device.mysql)
            sheet.cell(row=row, column=9, value=device.uname)
            sheet.cell(row=row, column=10, value=device.scanned)
            sheet.cell(row=row, column=11, value="; ".join(device.errors))

        # Auto-adjust column widths
        for column_cells in sheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(
                max_length + 2, 50
            )

        workbook.save(output_path)
        logger.info("Excel report generated at %s", output_path)
        return output_path

    def _generate_json_report(
        self, devices: list[Device], filename: str | None = None
    ) -> str:
        """Generate a JSON report.

        Args:
            devices: The list of devices to include in the report.
            filename: Optional custom filename (without extension).

        Returns:
            The path to the generated JSON file.
        """
        output_path = self._get_output_path("json", filename)

        # Create comprehensive report structure
        report_data = {
            "metadata": {
                "generated_at": datetime.now().isoformat(),
                "total_devices": len(devices),
                "alive_devices": sum(1 for d in devices if d.alive),
                "version": "1.0",
            },
            "devices": [device.to_dict() for device in devices],
            "summary": {
                "ssh_enabled": sum(1 for d in devices if d.ssh),
                "snmp_enabled": sum(1 for d in devices if d.snmp),
                "mysql_enabled": sum(1 for d in devices if d.mysql),
                "scanned": sum(1 for d in devices if d.scanned),
                "with_errors": sum(1 for d in devices if d.errors),
            },
        }

        with Path(output_path).open("w", encoding="utf-8") as f:
            json.dump(report_data, f, indent=2)

        logger.info("JSON report generated at %s", output_path)
        return output_path

    def _generate_pdf_report(
        self, devices: list[Device], filename: str | None = None
    ) -> str:
        """Generate a PDF report.

        Args:
            devices: The list of devices to include in the report.
            filename: Optional custom filename (without extension).

        Returns:
            The path to the generated PDF file.

        Raises:
            ValueError: If fpdf2 is not installed.
        """
        if not FPDF_AVAILABLE:
            raise ValueError("PDF export requires fpdf2. Install with: uv add fpdf2")

        from fpdf import FPDF  # noqa: PLC0415

        output_path = self._get_output_path("pdf", filename)

        # Create PDF document
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()

        # Title
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(0, 10, "Network Discovery Report", ln=True, align="C")

        # Metadata
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(
            0, 6, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", ln=True
        )
        pdf.cell(0, 6, f"Total Devices: {len(devices)}", ln=True)
        pdf.cell(0, 6, f"Alive Devices: {sum(1 for d in devices if d.alive)}", ln=True)
        pdf.ln(5)

        # Table header
        pdf.set_font("Helvetica", "B", 9)
        col_widths = [15, 35, 30, 15, 15, 15, 15, 50]
        headers = ["ID", "Host", "IP", "Alive", "SSH", "SNMP", "MySQL", "Errors"]

        for i, header in enumerate(headers):
            pdf.cell(col_widths[i], 8, header, border=1, align="C")
        pdf.ln()

        # Table data
        pdf.set_font("Helvetica", "", 8)
        for device in devices:
            # Check if we need a new page
            if pdf.get_y() > 270:
                pdf.add_page()
                pdf.set_font("Helvetica", "B", 9)
                for i, header in enumerate(headers):
                    pdf.cell(col_widths[i], 8, header, border=1, align="C")
                pdf.ln()
                pdf.set_font("Helvetica", "", 8)

            row_data = [
                str(device.id),
                device.host[:20] if len(device.host) > 20 else device.host,
                device.ip,
                "Yes" if device.alive else "No",
                "Yes" if device.ssh else "No",
                "Yes" if device.snmp else "No",
                "Yes" if device.mysql else "No",
                ("; ".join(device.errors))[:30] if device.errors else "",
            ]

            for i, data in enumerate(row_data):
                pdf.cell(col_widths[i], 7, data, border=1)
            pdf.ln()

        # Summary section
        pdf.ln(10)
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 8, "Summary", ln=True)
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(0, 6, f"SSH Enabled: {sum(1 for d in devices if d.ssh)}", ln=True)
        pdf.cell(0, 6, f"SNMP Enabled: {sum(1 for d in devices if d.snmp)}", ln=True)
        pdf.cell(0, 6, f"MySQL Enabled: {sum(1 for d in devices if d.mysql)}", ln=True)
        pdf.cell(
            0, 6, f"Devices with Errors: {sum(1 for d in devices if d.errors)}", ln=True
        )

        pdf.output(output_path)
        logger.info("PDF report generated at %s", output_path)
        return output_path

    def export_all_formats(
        self, devices: list[Device], filename: str | None = None
    ) -> dict[str, str]:
        """Export devices to all supported formats.

        Args:
            devices: The list of devices to include in the reports.
            filename: Optional base filename (without extension).

        Returns:
            Dictionary mapping format names to output file paths.
        """
        results = {}
        for format_type in self.supported_formats:
            try:
                results[format_type] = self.generate_report(
                    devices, format_type, filename
                )
            except Exception as e:
                logger.error("Failed to export %s: %s", format_type, e)
                results[format_type] = None
        return results
